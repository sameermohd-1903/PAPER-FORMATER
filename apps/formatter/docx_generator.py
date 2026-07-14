from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os
from pathlib import Path
from django.conf import settings
from .paper_formatter import PaperFormatter
class DocxGenerator(PaperFormatter):
    """DOCX generation using python-docx"""
    
    @classmethod
    def generate_paper(cls, paper):
        """Generate DOCX file for the paper"""
        # Create filename
        filename = f"paper_{paper.id}_{paper.created_at.strftime('%Y%m%d_%H%M%S')}.docx"
        docx_dir = Path(settings.MEDIA_ROOT) / 'generated_papers' / 'docx'
        docx_dir.mkdir(parents=True, exist_ok=True)
        filepath = str(docx_dir / filename)
        
        # Create document
        doc = Document()
        
        # Header
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title.add_run(paper.college_name.upper())
        title_run.bold = True
        title_run.font.size = Inches(0.2)
        
        doc.add_paragraph()  # Empty line
        
        # Subject
        subject = doc.add_paragraph()
        subject.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subject.add_run(f"Subject: {paper.subject}").bold = True
        
        # Time and marks
        details = doc.add_paragraph()
        details.alignment = WD_ALIGN_PARAGRAPH.CENTER
        details.add_run(f"Time: {paper.time_allowed}").bold = True
        details.add_run("\n")
        details.add_run(f"Maximum Marks: {paper.maximum_marks}").bold = True
        
        doc.add_paragraph()  # Empty line
        
        # Sections
        # Section A
        section_a_questions = list(paper.section_a_questions.all())
        if section_a_questions:
            section_a_title = doc.add_paragraph()
            section_a_title.add_run("SECTION A").bold = True
            section_a_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
            
            for i, question in enumerate(section_a_questions, 1):
                q_para = doc.add_paragraph()
                q_para.add_run(f"Q{i}. ").bold = True
                q_para.add_run(question.question_text)
                marks_para = doc.add_paragraph()
                marks_para.add_run(f"[Marks: {paper.pattern.section_a_marks}]").italic = True
                marks_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                doc.add_paragraph()  # Space between questions
        
        # Section B
        section_b_questions = list(paper.section_b_questions.all())
        if section_b_questions:
            section_b_title = doc.add_paragraph()
            section_b_title.add_run("SECTION B").bold = True
            section_b_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
            
            for i, question in enumerate(section_b_questions, 1):
                q_para = doc.add_paragraph()
                q_para.add_run(f"Q{i}. ").bold = True
                q_para.add_run(question.question_text)
                marks_para = doc.add_paragraph()
                marks_para.add_run(f"[Marks: {paper.pattern.section_b_marks}]").italic = True
                marks_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                doc.add_paragraph()  # Space between questions
        
        # Section C
        section_c_questions = list(paper.section_c_questions.all())
        if section_c_questions:
            section_c_title = doc.add_paragraph()
            section_c_title.add_run("SECTION C").bold = True
            section_c_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
            
            for i, question in enumerate(section_c_questions, 1):
                q_para = doc.add_paragraph()
                q_para.add_run(f"Q{i}. ").bold = True
                q_para.add_run(question.question_text)
                marks_para = doc.add_paragraph()
                marks_para.add_run(f"[Marks: {paper.pattern.section_c_marks}]").italic = True
                marks_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                doc.add_paragraph()  # Space between questions
        
        # Save document
        doc.save(filepath)
        
        return f"generated_papers/docx/{filename}"